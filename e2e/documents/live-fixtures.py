"""Synthetic, reproducible UI fixtures. Never uses user or production documents."""

import hashlib
import json
import math
import pathlib
import statistics
import sys
import zipfile

import xlsxwriter
from pptx import Presentation
from pptx.util import Inches

root = pathlib.Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=True)
count = int(sys.argv[2]) if len(sys.argv) > 2 else 300000
output = root / f"synthetic-{count}-rows.xlsx"
book = xlsxwriter.Workbook(output, {"constant_memory": True})
sheet = book.add_worksheet("Measurements")
sheet.write_row(0, 0, ["record_id", "measurement", "category"] + [f"noise_{i}" for i in range(10)])
values = []
groups = {}
passed = failed = missing = errors = 0
for i in range(1, count + 1):
    sheet.write_number(i, 0, i)
    group = f"Group-{i % 12:02d}"
    sheet.write_string(i, 2, group)
    groups[group] = groups.get(group, 0) + 1
    if i % 1000 == 0:
        missing += 1
    elif i % 1000 == 1:
        sheet.write_formula(i, 1, "=1/0", None, "#DIV/0!")
        errors += 1
    else:
        value = 10.0 if i % 50 == 0 else round(9 + (i % 101) / 50, 2)
        if i % 50 == 0:
            sheet.write_formula(i, 1, "=5+5", None, value)
        else:
            sheet.write_number(i, 1, value)
        values.append(value)
        if 9.4 <= value <= 10.6:
            passed += 1
        else:
            failed += 1
    for column in range(3, 13):
        sheet.write_string(i, column, hashlib.sha256(f"pr138:20261015:{i}:{column}".encode()).hexdigest())
book.close()
with zipfile.ZipFile(output) as archive:
    expanded = sum(item.file_size for item in archive.infolist())
golden = {
    "file": output.name, "rows": count, "bytes": output.stat().st_size,
    "expandedBytes": expanded, "columns": 13, "selectedColumns": ["A", "B", "C"],
    "bounds": {"lower": 9.4, "upper": 10.6},
    "classification": {"pass": passed, "fail": failed, "incomplete": missing + errors},
    "validCount": len(values), "missingCount": missing, "errorCount": errors,
    "mean": statistics.fmean(values), "min": min(values), "max": max(values),
    "sampleStddev": statistics.stdev(values), "distinctRecordIds": count, "groups": groups,
}
(root / "golden.json").write_text(json.dumps(golden, indent=2), encoding="utf-8")

small = xlsxwriter.Workbook(root / "small-baseline.xlsx")
small.add_worksheet("Data").write_row(0, 0, ["name", "value"])
small.worksheets()[0].write_column(1, 0, ["alpha", "beta", "gamma"])
small.worksheets()[0].write_column(1, 1, [8, 10, 12])
small.close()

deck = Presentation()
for title, body in [
    ("Synthetic regression deck", "PR138 / PR144 / PR146 test environment only"),
    ("Known measurements", "alpha: 8\nbeta: 10\ngamma: 12\nmean: 10"),
    ("Acceptance", "Read all three slides; create a new deck without overwriting this input."),
]:
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
deck.save(root / "baseline-three-slides.pptx")
(root / "small-baseline.txt").write_text("Synthetic local-file regression. Expected marker: PR138_LOCAL_FILE_OK.\n", encoding="utf-8")
print(json.dumps(golden, ensure_ascii=False))
