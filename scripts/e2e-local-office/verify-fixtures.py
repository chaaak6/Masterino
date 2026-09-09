#!/usr/bin/env python3
"""Independently reread QA inputs; this is not product acceptance."""
import hashlib
import json
import sys
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook

root = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/masterino-office-qa')
oracle = json.loads((root/'oracle.json').read_text())
for name, digest in oracle['sha256'].items():
    assert hashlib.sha256((root/name).read_bytes()).hexdigest() == digest, name
for name, expected in oracle['workbooks'].items():
    wb = load_workbook(root/name, read_only=True, data_only=True)
    total = count = 0
    regions, months = {}, {}
    for row in wb['Sales'].iter_rows(min_row=2, values_only=True):
        count += 1
        total += row[6]
        regions[row[2]] = regions.get(row[2], 0) + row[6]
        months[row[1]] = months.get(row[1], 0) + row[6]
    assert (count, total, regions, months) == (expected['rows'], expected['revenue'], expected['byRegion'], expected['byMonth']), name
    wb.close()
assert load_workbook(root/'formula-cache.xlsx', data_only=True).active['A2'].value == 999
assert load_workbook(root/'formula-cache.xlsx', data_only=False).active['A2'].value == '=1+2'
with ZipFile(root/'reordered.pptx') as z:
    relations = {r.attrib['Id']: r.attrib['Target'] for r in ET.fromstring(z.read('ppt/_rels/presentation.xml.rels'))}
    presentation = ET.fromstring(z.read('ppt/presentation.xml'))
    order = []
    for slide in presentation.find('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst'):
        target = relations[slide.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']]
        order.append(''.join(ET.fromstring(z.read('ppt/'+target)).itertext()))
    assert order == oracle['presentationOrder']
print('Fixture hashes, workbook aggregates, formula cache and relationship-based slide order verified. Product E2E remains separate.')
