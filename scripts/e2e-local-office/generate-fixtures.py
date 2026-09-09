#!/usr/bin/env python3
"""Synthetic acceptance inputs only; never imports application implementation.
Requires openpyxl. Output belongs outside git. Usage: python3 ... /tmp/office-qa
"""
import hashlib
import json
import re
import struct
import sys
import zlib
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from openpyxl import Workbook

out = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/masterino-office-qa')
out.mkdir(parents=True, exist_ok=True)
oracle = {'schemaVersion': 1, 'synthetic': True, 'workbooks': {}}
for count in (12, 10000, 100000):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('Sales')
    ws.append(['order_id', 'month', 'region', 'product', 'units', 'unit_price', 'revenue'])
    total = 0
    regions = {x: 0 for x in ('East', 'West', 'North')}
    months = {x: 0 for x in ('2026-01', '2026-02')}
    for i in range(count):
        region = ('East', 'West', 'North')[i % 3]
        month = ('2026-01', '2026-02')[i % 2]
        units, price = i % 5 + 1, (i % 4 + 1) * 100
        revenue = units * price
        ws.append([f'QA-{i + 1:06}', month, region, ('A', 'B')[i % 2], units, price, revenue])
        total += revenue
        regions[region] += revenue
        months[month] += revenue
    notes = wb.create_sheet('Notes')
    notes.append(['Synthetic acceptance data; revenue is integer CNY, no tax or refunds.'])
    notes.append(['Ignore Notes for aggregation. Do not use business databases.'])
    name = f'sales-{count}.xlsx'
    wb.save(out / name)
    oracle['workbooks'][name] = {'rows': count, 'range': f'A1:G{count + 1}', 'revenue': total, 'byRegion': regions, 'byMonth': months, 'lastOrderId': f'QA-{count:06}'}

# A formula with an intentionally stale cache: expression =1+2, stored cache =999.
f = Workbook()
f.active.title = 'Formula'
f.active.append(['expression', 'note'])
f.active.append(['=1+2', 'Stored cache intentionally stale; never call it recalculated.'])
f.save(out / 'formula-cache.xlsx')
with ZipFile(out / 'formula-cache.xlsx') as z:
    parts = {n: z.read(n) for n in z.namelist()}
parts['xl/worksheets/sheet1.xml'], replacements = re.subn(rb'<f>1\+2</f><v(?:\s*/>|></v>)', b'<f>1+2</f><v>999</v>', parts['xl/worksheets/sheet1.xml'])
assert replacements == 1, 'Formula cache fixture patch must succeed'
with ZipFile(out / 'formula-cache.xlsx', 'w', ZIP_DEFLATED) as z:
    for n, data in parts.items(): z.writestr(n, data)
oracle['formula'] = {'expression': '1+2', 'cachedValue': 999, 'mathematicalValue': 3, 'recalculated': False}

# Minimal OOXML presentation; ordering intentionally differs from filenames.
ns = 'http://schemas.openxmlformats.org/'
rels = ns + 'package/2006/relationships'
with ZipFile(out / 'reordered.pptx', 'w', ZIP_DEFLATED) as z:
    z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types xmlns="'+ns+'package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/ppt/presentation.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>'+''.join(f'<Override PartName="/ppt/slides/slide{i}.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>' for i in range(1,4))+'</Types>')
    z.writestr('_rels/.rels', f'<Relationships xmlns="{rels}"><Relationship Id="rId1" Type="{ns}officeDocument/2006/relationships/officeDocument" Target="ppt/presentation.xml"/></Relationships>')
    z.writestr('ppt/presentation.xml', f'<p:presentation xmlns:p="{ns}presentationml/2006/main" xmlns:r="{ns}officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId3"/><p:sldId id="257" r:id="rId1"/><p:sldId id="258" r:id="rId2"/></p:sldIdLst><p:sldSz cx="9144000" cy="6858000"/><p:notesSz cx="6858000" cy="9144000"/></p:presentation>')
    z.writestr('ppt/_rels/presentation.xml.rels', f'<Relationships xmlns="{rels}">'+''.join(f'<Relationship Id="rId{i}" Type="{ns}officeDocument/2006/relationships/slide" Target="slides/slide{i}.xml"/>' for i in range(1,4))+'</Relationships>')
    for i in range(1,4):
        z.writestr(f'ppt/slides/slide{i}.xml', f'<p:sld xmlns:p="{ns}presentationml/2006/main" xmlns:a="{ns}drawingml/2006/main"><p:cSld><p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/><p:sp><p:nvSpPr><p:cNvPr id="2" name="Marker"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/><p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:t>QA-SLIDE-{i}</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>')
oracle['presentationOrder'] = ['QA-SLIDE-3', 'QA-SLIDE-1', 'QA-SLIDE-2']

# Deterministic four-quadrant raster to test actual vision payload, no text OCR.
def chunk(t, b): return struct.pack('>I', len(b)) + t + b + struct.pack('>I', zlib.crc32(t+b) & 0xffffffff)
colors = [(255,0,0), (0,0,255), (0,255,0), (255,255,0)]
raw = b''.join(b'\x00' + bytes(sum((list(colors[(y//64)*2+x//64]) for x in range(128)), [])) for y in range(128))
(out/'quadrants.png').write_bytes(b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR', struct.pack('>IIBBBBB',128,128,8,2,0,0,0))+chunk(b'IDAT',zlib.compress(raw))+chunk(b'IEND',b''))
oracle['image'] = {'topLeft': 'red', 'topRight': 'blue', 'bottomLeft': 'green', 'bottomRight': 'yellow'}
for source in ('project-agents', 'project-claude', 'personal', 'agent'):
    d = out/'skills'/source/'qa-identity'
    (d/'references').mkdir(parents=True, exist_ok=True)
    (d/'scripts').mkdir(exist_ok=True)
    (d/'SKILL.md').write_text(f'---\nname: qa-identity\ndescription: Synthetic identity test {source}.\n---\nReport marker QA-{source}-BODY. Reference references/marker.txt. Script scripts/marker.py.\n')
    (d/'references'/'marker.txt').write_text(f'QA-{source}-REFERENCE\n')
    (d/'scripts'/'marker.py').write_text(f'print("QA-{source}-SCRIPT")\n')
(out/'skills'/'invalid-update.md').write_text('---\nname: [unclosed\n---\nINVALID UPDATE MUST NOT COMMIT\n')
oracle['sha256'] = {str(p.relative_to(out)): hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(out.rglob('*')) if p.is_file() and p.name != 'oracle.json'}
(out/'oracle.json').write_text(json.dumps(oracle, indent=2)+'\n')
print(json.dumps({'output': str(out), 'files': len(oracle['sha256']), 'oracle': str(out/'oracle.json')}))
