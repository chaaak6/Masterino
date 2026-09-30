"""Acceptance test for a relocated, general-purpose bundled interpreter."""
import hashlib
import importlib.metadata
import json
import os
import pathlib
import subprocess
import sys
import tempfile
import unittest
import zipfile

from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


class RuntimeTests(unittest.TestCase):
    def test_runtime_and_preinstalled_packages(self):
        runtime = pathlib.Path(sys.executable).resolve().parent
        if runtime.name == "bin":
            runtime = runtime.parent
        manifest = json.loads((runtime / "manifest.json").read_text())
        self.assertEqual(sys.version.split()[0], manifest["version"])
        for name, version in manifest["packages"].items():
            self.assertEqual(importlib.metadata.version(name), version)

    def test_isolated_scripts_ignore_host_and_workspace_packages(self):
        with tempfile.TemporaryDirectory(prefix="masterino-isolation-") as directory:
            root = pathlib.Path(directory)
            (root / "pptx.py").write_text("raise RuntimeError('workspace shadow imported')")
            script = root / "check.py"
            script.write_text("import json,pptx; print(json.dumps(pptx.__file__))")
            environment = {**os.environ, "PYTHONPATH": str(root), "PYTHONUSERBASE": str(root),
                           "PYTHONHOME": str(root / "missing-home")}
            user_site = subprocess.check_output(
                [sys.executable, "-B", "-X", "utf8", "-c",
                 "import site; print(site.getusersitepackages())"],
                env={**os.environ, "PYTHONHOME": "", "PYTHONPATH": "", "PYTHONUSERBASE": str(root)},
                text=True, encoding="utf8", timeout=30,
            ).strip()
            pathlib.Path(user_site).mkdir(parents=True, exist_ok=True)
            (pathlib.Path(user_site) / "pptx.py").write_text("raise RuntimeError('user site shadow imported')")
            output = subprocess.check_output(
                [sys.executable, "-I", "-B", "-X", "utf8", str(script)],
                env=environment, text=True, encoding="utf8", timeout=30,
            )
            self.assertEqual(pathlib.Path(json.loads(output)).resolve(),
                             pathlib.Path(sys.modules["pptx"].__file__).resolve())

    def test_additional_packages_stay_in_virtual_environment(self):
        runtime = pathlib.Path(sys.executable).resolve().parent
        if runtime.name == "bin":
            runtime = runtime.parent

        def runtime_digest():
            digest = hashlib.sha256()
            for file in sorted(runtime.rglob("*")):
                if file.is_file():
                    digest.update(str(file.relative_to(runtime)).encode())
                    digest.update(file.read_bytes())
            return digest.hexdigest()

        before = runtime_digest()
        with tempfile.TemporaryDirectory(prefix="masterino-dependencies-") as directory:
            root = pathlib.Path(directory)
            environment = root / "项目 环境"
            subprocess.run([sys.executable, "-I", "-B", "-X", "utf8", "-m", "venv",
                            "--without-pip", "--system-site-packages", str(environment)],
                           check=True, timeout=60)
            python = environment / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
            wheel = root / "masterino_test_addon-1.0-py3-none-any.whl"
            info = "masterino_test_addon-1.0.dist-info/"
            with zipfile.ZipFile(wheel, "w") as archive:
                archive.writestr("masterino_test_addon/__init__.py", "MARKER = 'venv-only'\n")
                archive.writestr(info + "METADATA", "Metadata-Version: 2.1\nName: masterino-test-addon\nVersion: 1.0\n")
                archive.writestr(info + "WHEEL", "Wheel-Version: 1.0\nRoot-Is-Purelib: true\nTag: py3-none-any\n")
                archive.writestr(info + "RECORD", "")
            subprocess.run([str(python), "-I", "-B", "-X", "utf8", "-m", "pip", "install",
                            "--no-index", "--no-deps", "--no-compile", str(wheel)],
                           check=True, capture_output=True, timeout=60)
            code = ("import json,pptx,masterino_test_addon as addon; "
                    "print(json.dumps([addon.MARKER,addon.__file__,pptx.__file__]))")
            marker, installed, pptx_file = json.loads(subprocess.check_output(
                [str(python), "-I", "-B", "-X", "utf8", "-c", code],
                text=True, encoding="utf8", timeout=30))
            self.assertEqual(marker, "venv-only")
            self.assertTrue(pathlib.Path(installed).is_relative_to(environment))
            self.assertEqual(pathlib.Path(pptx_file).resolve(),
                             pathlib.Path(sys.modules["pptx"].__file__).resolve())
            self.assertFalse((runtime / "lib/python3.12/site-packages/masterino_test_addon").exists())
        self.assertEqual(runtime_digest(), before, "Bundled runtime was modified")

    def test_create_inspect_and_edit_rich_presentation(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            image = root / "image.png"
            Image.new("RGB", (120, 80), "#2563eb").save(image)
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            text = group.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            text.text = "中文组合标题"
            table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(3), Inches(2)).table
            table.cell(0, 0).text = "收入"
            table.cell(1, 0).text = "华东"
            table.cell(1, 1).text = "120"
            data = CategoryChartData()
            data.categories = ["一月", "二月"]
            data.add_series("销售", [100, 120])
            chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1), Inches(4), Inches(3), data).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = "销售趋势"
            slide.shapes.add_picture(str(image), Inches(1), Inches(5))
            slide.notes_slide.notes_text_frame.text = "中文演讲备注"
            original = root / "中文课件.pptx"
            edited = root / "修改课件.pptx"
            deck.save(original)
            loaded = Presentation(original)
            self.assertEqual(loaded.slides[0].shapes[0].shapes[0].text, "中文组合标题")
            loaded.slides[0].shapes[0].shapes[0].text = "修改后的标题"
            loaded.save(edited)
            result = Presentation(edited)
            self.assertEqual(result.slides[0].shapes[0].shapes[0].text, "修改后的标题")
            self.assertEqual(result.slides[0].shapes[1].table.cell(1, 1).text, "120")
            self.assertEqual(result.slides[0].shapes[2].chart.chart_title.text_frame.text, "销售趋势")
            self.assertIn("中文演讲备注", result.slides[0].notes_slide.notes_text_frame.text)
            with zipfile.ZipFile(original) as before, zipfile.ZipFile(edited) as after:
                self.assertEqual(set(before.namelist()), set(after.namelist()))
                for name in before.namelist():
                    if name != "ppt/slides/slide1.xml":
                        self.assertEqual(before.read(name), after.read(name), name)
                etree.fromstring(after.read("ppt/slides/slide1.xml"))


if __name__ == "__main__":
    unittest.main()
