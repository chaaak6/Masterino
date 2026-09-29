"""Fixed-command PowerPoint worker used by the Electron main process."""

import hashlib
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


PROTOCOL_VERSION = 1
MAX_INPUT_BYTES = 100 * 1024 * 1024
MAX_OPERATIONS = 200
MAX_TEXT_CHARS = 10_000


class WorkerError(Exception):
    pass


def require(condition, code):
    if not condition:
        raise WorkerError(code)


def digest(path):
    sha = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            sha.update(chunk)
    return sha.hexdigest()


def source_path(request):
    raw = request.get("path")
    require(isinstance(raw, str) and Path(raw).is_absolute(), "PRESENTATION_INVALID_PATH")
    source = Path(raw)
    require(source.suffix.lower() == ".pptx", "PRESENTATION_INVALID_PATH")
    require(source.is_file(), "PRESENTATION_SOURCE_MISSING")
    require(source.stat().st_size <= MAX_INPUT_BYTES, "PRESENTATION_SOURCE_TOO_LARGE")
    return source


def frame(shape):
    return {
        "x": round(shape.left / 914400, 4),
        "y": round(shape.top / 914400, 4),
        "w": round(shape.width / 914400, 4),
        "h": round(shape.height / 914400, 4),
    }


def inspect(request):
    source = source_path(request)
    expected = request.get("expectedSha256")
    sha = digest(source)
    if expected is not None:
        require(expected == sha, "PRESENTATION_SOURCE_CHANGED")
    presentation = Presentation(source)
    indices = request.get("slideIndices")
    if indices is None:
        indices = list(range(1, min(len(presentation.slides), 20) + 1))
    require(
        isinstance(indices, list)
        and len(indices) <= 20
        and all(type(index) is int and 1 <= index <= len(presentation.slides) for index in indices)
        and len(set(indices)) == len(indices),
        "PRESENTATION_INVALID_SLIDE_SELECTION",
    )
    slides = []
    shape_count = 0
    text_chars = 0
    for index in indices:
        slide = presentation.slides[index - 1]
        shapes = []
        for shape in slide.shapes:
            shape_count += 1
            require(shape_count <= 500, "PRESENTATION_INSPECTION_LIMIT")
            paragraphs = []
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    runs = []
                    for run in paragraph.runs:
                        text_chars += len(run.text)
                        require(text_chars <= 64_000, "PRESENTATION_INSPECTION_LIMIT")
                        runs.append(run.text)
                    paragraphs.append({"runs": runs})
            shapes.append(
                {
                    "shapeId": shape.shape_id,
                    "name": shape.name,
                    "type": str(shape.shape_type),
                    "frame": frame(shape),
                    "paragraphs": paragraphs,
                }
            )
        slides.append({"slideIndex": index, "shapes": shapes})
    return {
        "protocolVersion": PROTOCOL_VERSION,
        "sha256": sha,
        "totalSlides": len(presentation.slides),
        "hasMore": len(indices) < len(presentation.slides),
        "slides": slides,
    }


def get_slide(presentation, operation):
    index = operation.get("slideIndex")
    require(type(index) is int and 1 <= index <= len(presentation.slides), "PRESENTATION_SLIDE_NOT_FOUND")
    return presentation.slides[index - 1]


def get_shape(slide, operation):
    shape_id = operation.get("shapeId")
    require(type(shape_id) is int, "PRESENTATION_SHAPE_NOT_FOUND")
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise WorkerError("PRESENTATION_SHAPE_NOT_FOUND")


def get_frame(value):
    require(isinstance(value, dict), "PRESENTATION_INVALID_FRAME")
    coordinates = [value.get(key) for key in ("x", "y", "w", "h")]
    require(
        all(type(item) in (int, float) and 0 <= item <= 100 for item in coordinates)
        and coordinates[2] > 0
        and coordinates[3] > 0,
        "PRESENTATION_INVALID_FRAME",
    )
    return [Inches(item) for item in coordinates]


def apply_operation(presentation, operation):
    require(isinstance(operation, dict), "PRESENTATION_INVALID_OPERATION")
    slide = get_slide(presentation, operation)
    kind = operation.get("op")
    if kind == "replaceTextRun":
        shape = get_shape(slide, operation)
        require(shape.has_text_frame, "PRESENTATION_TEXT_NOT_FOUND")
        paragraph_index = operation.get("paragraphIndex")
        run_index = operation.get("runIndex")
        text = operation.get("text")
        require(
            type(paragraph_index) is int
            and 0 <= paragraph_index < len(shape.text_frame.paragraphs),
            "PRESENTATION_TEXT_NOT_FOUND",
        )
        runs = shape.text_frame.paragraphs[paragraph_index].runs
        require(type(run_index) is int and 0 <= run_index < len(runs), "PRESENTATION_TEXT_NOT_FOUND")
        require(isinstance(text, str) and len(text) <= MAX_TEXT_CHARS, "PRESENTATION_INVALID_TEXT")
        runs[run_index].text = text
    elif kind == "setShapeFrame":
        shape = get_shape(slide, operation)
        shape.left, shape.top, shape.width, shape.height = get_frame(operation.get("frame"))
    elif kind == "addTextBox":
        text = operation.get("text")
        require(isinstance(text, str) and 0 < len(text) <= MAX_TEXT_CHARS, "PRESENTATION_INVALID_TEXT")
        shape = slide.shapes.add_textbox(*get_frame(operation.get("frame")))
        shape.text = text
    else:
        raise WorkerError("PRESENTATION_UNSUPPORTED_OPERATION")


def apply(request):
    source = source_path(request)
    output_raw = request.get("outputPath")
    require(isinstance(output_raw, str) and Path(output_raw).is_absolute(), "PRESENTATION_INVALID_PATH")
    output = Path(output_raw)
    require(output.suffix.lower() == ".pptx", "PRESENTATION_INVALID_PATH")
    require(output != source and not output.exists(), "PRESENTATION_OUTPUT_EXISTS")
    expected = request.get("expectedSha256")
    require(isinstance(expected, str) and re.fullmatch(r"[0-9a-f]{64}", expected), "PRESENTATION_EXPECTED_HASH_REQUIRED")
    require(digest(source) == expected, "PRESENTATION_SOURCE_CHANGED")
    operations = request.get("operations")
    require(isinstance(operations, list) and 1 <= len(operations) <= MAX_OPERATIONS, "PRESENTATION_INVALID_OPERATIONS")
    presentation = Presentation(source)
    for operation in operations:
        apply_operation(presentation, operation)
    presentation.save(output)
    return {
        "protocolVersion": PROTOCOL_VERSION,
        "appliedOperations": len(operations),
        "sha256": digest(output),
        "slides": len(presentation.slides),
    }


def main():
    try:
        command = sys.argv[1] if len(sys.argv) == 2 else ""
        if command == "health":
            result = {"protocolVersion": PROTOCOL_VERSION, "pythonPptx": True}
        else:
            request = json.load(sys.stdin)
            require(isinstance(request, dict), "PRESENTATION_INVALID_REQUEST")
            if command == "inspect":
                result = inspect(request)
            elif command == "apply":
                result = apply(request)
            else:
                raise WorkerError("PRESENTATION_INVALID_COMMAND")
        print(json.dumps(result))
    except (WorkerError, OSError, ValueError, KeyError) as error:
        code = str(error) if isinstance(error, WorkerError) else "PRESENTATION_WORKER_FAILED"
        print(json.dumps({"error": code}))
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
