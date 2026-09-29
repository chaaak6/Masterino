import hashlib
import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


WORKER = Path(__file__).with_name("worker.py")


def invoke(command, payload):
    executable = os.environ.get("PPT_WORKER_BINARY")
    completed = subprocess.run(
        [executable, command] if executable else [sys.executable, str(WORKER), command],
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
        input=json.dumps(payload),
        text=True,
        encoding="utf-8",
        capture_output=True,
        timeout=20,
        check=False,
    )
    return completed.returncode, json.loads(completed.stdout)


class WorkerContractTest(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.addCleanup(self.folder.cleanup)
        self.source = Path(self.folder.name) / "source.pptx"
        self.output = Path(self.folder.name) / "output.pptx"
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        first = shape.text_frame.paragraphs[0].add_run()
        first.text = "Old title"
        first.font.bold = True
        second = shape.text_frame.paragraphs[0].add_run()
        second.text = " stays"
        second.font.italic = True
        deck.save(self.source)
        self.shape_id = shape.shape_id
        self.digest = hashlib.sha256(self.source.read_bytes()).hexdigest()

    def test_inspect_and_edit_preserve_source_and_other_runs(self):
        code, inspected = invoke("inspect", {"path": str(self.source)})
        self.assertEqual(code, 0)
        self.assertEqual(inspected["sha256"], self.digest)
        self.assertEqual(inspected["totalSlides"], 1)
        self.assertEqual(inspected["slides"][0]["shapes"][0]["shapeId"], self.shape_id)
        self.assertEqual(
            inspected["slides"][0]["shapes"][0]["paragraphs"][0]["runs"][0],
            "Old title",
        )

        code, edited = invoke(
            "apply",
            {
                "path": str(self.source),
                "outputPath": str(self.output),
                "expectedSha256": self.digest,
                "operations": [
                    {
                        "op": "replaceTextRun",
                        "slideIndex": 1,
                        "shapeId": self.shape_id,
                        "paragraphIndex": 0,
                        "runIndex": 0,
                        "text": "New title",
                    }
                ],
            },
        )
        self.assertEqual(code, 0, edited)
        self.assertEqual(edited["appliedOperations"], 1)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.digest)
        changed = Presentation(self.output).slides[0].shapes[0].text_frame.paragraphs[0].runs
        self.assertEqual([run.text for run in changed], ["New title", " stays"])
        self.assertTrue(changed[0].font.bold)
        self.assertTrue(changed[1].font.italic)

    def test_rejects_stale_source_without_output(self):
        code, result = invoke(
            "apply",
            {
                "path": str(self.source),
                "outputPath": str(self.output),
                "expectedSha256": "0" * 64,
                "operations": [],
            },
        )
        self.assertNotEqual(code, 0)
        self.assertEqual(result["error"], "PRESENTATION_SOURCE_CHANGED")
        self.assertFalse(self.output.exists())

    def test_frame_and_textbox_edits_do_not_change_original(self):
        code, result = invoke(
            "apply",
            {
                "path": str(self.source),
                "outputPath": str(self.output),
                "expectedSha256": self.digest,
                "operations": [
                    {
                        "op": "setShapeFrame",
                        "slideIndex": 1,
                        "shapeId": self.shape_id,
                        "frame": {"x": 2, "y": 1, "w": 3, "h": 1},
                    },
                    {
                        "op": "addTextBox",
                        "slideIndex": 1,
                        "frame": {"x": 1, "y": 3, "w": 4, "h": 1},
                        "text": "新备注",
                    },
                ],
            },
        )
        self.assertEqual(code, 0, result)
        changed = Presentation(self.output).slides[0]
        self.assertEqual(round(changed.shapes[0].left / 914400), 2)
        self.assertEqual(changed.shapes[1].text, "新备注")
        code, inspected = invoke("inspect", {"path": str(self.output)})
        self.assertEqual(code, 0, inspected)
        self.assertEqual(inspected["slides"][0]["shapes"][1]["paragraphs"][0]["runs"], ["新备注"])
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.digest)

    def test_rejects_unknown_shape_without_output(self):
        code, result = invoke(
            "apply",
            {
                "path": str(self.source),
                "outputPath": str(self.output),
                "expectedSha256": self.digest,
                "operations": [{"op": "setShapeFrame", "slideIndex": 1, "shapeId": 9999,
                                "frame": {"x": 2, "y": 1, "w": 3, "h": 1}}],
            },
        )
        self.assertNotEqual(code, 0)
        self.assertEqual(result["error"], "PRESENTATION_SHAPE_NOT_FOUND")
        self.assertFalse(self.output.exists())

    def test_rejects_non_pptx_output(self):
        code, result = invoke(
            "apply",
            {"path": str(self.source), "outputPath": str(self.output.with_suffix(".docx")),
             "expectedSha256": self.digest, "operations": [{"op": "addTextBox", "slideIndex": 1,
             "frame": {"x": 1, "y": 2, "w": 3, "h": 1}, "text": "No"}]},
        )
        self.assertNotEqual(code, 0)
        self.assertEqual(result["error"], "PRESENTATION_INVALID_PATH")


if __name__ == "__main__":
    unittest.main()
